import json
import os
import re
from fractions import Fraction
from typing import Optional, Dict, Any

import openpyxl

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

# Excel文件路径
excel_path = "/Users/bytedance/Documents/ai图/13x6雪花网-TT6-1057#/13x6雪花网-TT6-1057#.xlsx"

# 输出目录
output_dir = os.path.join(SCRIPT_DIR, "导入json数据")

# 确保目录存在
os.makedirs(output_dir, exist_ok=True)


UNICODE_FRACTIONS = {
    "¼": Fraction(1, 4),
    "½": Fraction(1, 2),
    "¾": Fraction(3, 4),
}

DML_KEYS = ("D", "M", "L")
EPSILON = 1e-9


def parse_fraction_number(text: str) -> Optional[float]:
    text = text.strip()
    if not text:
        return None

    normalized = (
        text.replace("／", "/")
        .replace("．", ".")
        .replace("，", ",")
        .replace("−", "-")
        .replace("–", "-")
    )

    try:
        return float(normalized.replace(",", ""))
    except (ValueError, TypeError):
        pass

    for char, fraction in UNICODE_FRACTIONS.items():
        if char in normalized:
            normalized = normalized.replace(
                char, f" {fraction.numerator}/{fraction.denominator}"
            )

    normalized = re.sub(r"\s+", " ", normalized).strip()

    mixed_match = re.fullmatch(r"(-?\d+)\s+(\d+/\d+)", normalized)
    if mixed_match:
        whole = int(mixed_match.group(1))
        frac = float(Fraction(mixed_match.group(2)))
        return whole + frac if whole >= 0 else whole - frac

    if re.fullmatch(r"-?\d+/\d+", normalized):
        return float(Fraction(normalized))

    return None


def to_num(val) -> Optional[float]:
    if val is None:
        return None

    if isinstance(val, (int, float)):
        return float(val)

    return parse_fraction_number(str(val))


def to_str(val) -> str:
    if val is None:
        return ""
    return str(val).strip()


def normalize_text(val) -> str:
    text = to_str(val)
    text = text.replace("（", "(").replace("）", ")")
    text = re.sub(r'[\s"“”\'‘’`]', "", text)
    return text


def strip_note_key(key: str, value: str) -> str:
    value = to_str(value)
    if not value:
        return ""

    normalized_key = normalize_text(key)
    normalized_value = normalize_text(value)
    if not normalized_key or normalized_key not in normalized_value:
        return value

    patterns = [
        rf"^\s*{re.escape(key)}\s*[：:]\s*",
        rf"^\s*{re.escape(key)}\s*",
    ]
    cleaned = value
    for pattern in patterns:
        cleaned = re.sub(pattern, "", cleaned, count=1)

    return cleaned.strip()


def clean_meta_value(value: str) -> str:
    value = to_str(value)
    if value in {"0", "None", "nan"}:
        return ""
    return value


def has_effective_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and abs(float(value)) > EPSILON


def get_present_dml_keys(dml_data: Dict[str, Any]) -> list[str]:
    return [key for key in DML_KEYS if has_effective_number(dml_data.get(key))]


def build_ratio_from_weight_data(weight_data: Dict[str, Any]) -> Optional[Dict[str, float]]:
    present_keys = get_present_dml_keys(weight_data)
    if "D" not in present_keys or len(present_keys) <= 1:
        return None

    total = sum(float(weight_data[key]) for key in present_keys)
    if total <= EPSILON:
        return None

    ratio: Dict[str, float] = {}
    for key in present_keys:
        normalized = round(float(weight_data[key]) / total * 10, 1)
        ratio[key] = int(normalized) if float(normalized).is_integer() else normalized
    return ratio


def build_default_ratio_from_keys(weight_keys: set[str]) -> Optional[Dict[str, float]]:
    if "D" not in weight_keys:
        return None

    keys = ["D"]
    if "M" in weight_keys:
        keys.append("M")
    if "L" in weight_keys:
        keys.append("L")

    if len(keys) <= 1:
        return None

    return {key: 1 for key in keys}


def analyze_imported_specs(
    machine_spec_list: list[Dict[str, Any]], manual_spec_list: list[Dict[str, Any]]
) -> Dict[str, Any]:
    size_keys: set[str] = set()
    weight_keys: set[str] = set()
    inferred_ratio: Optional[Dict[str, float]] = None

    for spec in machine_spec_list:
        size_data = spec.get("双针", {}).get("尺数", {})
        size_keys.update(get_present_dml_keys(size_data))

        for item in spec.get("裁断与重量", []):
            weight_data = item.get("重量g", {})
            weight_keys.update(get_present_dml_keys(weight_data))
            if inferred_ratio is None:
                inferred_ratio = build_ratio_from_weight_data(weight_data)

    for spec in manual_spec_list:
        for item in spec.get("裁断与重量", []):
            weight_data = item.get("重量g", {})
            weight_keys.update(get_present_dml_keys(weight_data))
            if inferred_ratio is None:
                inferred_ratio = build_ratio_from_weight_data(weight_data)

    has_only_d_size = "D" in size_keys and "M" not in size_keys and "L" not in size_keys
    has_extra_weight = "M" in weight_keys or "L" in weight_keys
    has_only_d_weight = "D" in weight_keys and not has_extra_weight

    if has_only_d_size and has_only_d_weight:
        wig_type = "纯色"
    elif has_only_d_size and "D" in weight_keys:
        wig_type = "间色"
    else:
        wig_type = "上下分"

    return {
        "wig_type": wig_type,
        "size_keys": size_keys,
        "weight_keys": weight_keys,
        "dml_ratio": inferred_ratio or build_default_ratio_from_keys(weight_keys),
    }


def parse_excel(path: str) -> Dict[str, Any]:
    """解析Excel文件，返回机器规格清单、人工规格清单和工艺说明"""
    wb = openpyxl.load_workbook(path, data_only=True)

    # 不能依赖 wb.active：有些 Excel 的激活 sheet 不是“规格书”页。
    # 这里遍历所有工作表，找到包含「分类」表头的那个 sheet 再解析。
    ws = None
    merged_value_map: Dict[tuple[int, int], Any] = {}

    def build_merged_map(sheet) -> Dict[tuple[int, int], Any]:
        m: Dict[tuple[int, int], Any] = {}
        for merged_range in sheet.merged_cells.ranges:
            top_left_value = sheet.cell(merged_range.min_row, merged_range.min_col).value
            for rr in range(merged_range.min_row, merged_range.max_row + 1):
                for cc in range(merged_range.min_col, merged_range.max_col + 1):
                    m[(rr, cc)] = top_left_value
        return m

    def get_display_value_factory(sheet, merged_map):
        def _get_display_value(row: int, col: int):
            value = sheet.cell(row, col).value
            if value is None or str(value).strip() == "":
                return merged_map.get((row, col))
            return value
        return _get_display_value

    for sheet in wb.worksheets:
        merged_map_try = build_merged_map(sheet)
        getv = get_display_value_factory(sheet, merged_map_try)
        found = False
        for r in range(1, sheet.max_row + 1):
            for c in range(1, sheet.max_column + 1):
                if normalize_text(getv(r, c)) == "分类":
                    found = True
                    break
            if found:
                break
        if found:
            ws = sheet
            merged_value_map = merged_map_try
            break

    if ws is None:
        raise ValueError("未找到「分类」表头")

    get_display_value = get_display_value_factory(ws, merged_value_map)

    def find_label_positions(*labels: str):
        normalized_labels = {normalize_text(label) for label in labels}
        positions = []
        for row in range(1, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                if normalize_text(get_display_value(row, col)) in normalized_labels:
                    positions.append((row, col))
        return positions

    def find_value_by_label(*labels: str) -> str:
        normalized_labels = {normalize_text(label) for label in labels}
        for row, col in find_label_positions(*labels):
            for next_col in range(col + 1, min(col + 4, ws.max_column) + 1):
                value = to_str(get_display_value(row, next_col))
                normalized_value = normalize_text(value)
                if (
                    value
                    and normalized_value not in normalized_labels
                    and not value.endswith(":")
                    and not value.endswith("：")
                ):
                    return value

        return ""

    # 根据表头文本定位分类列和主表头行（从 A 列开始扫描整张表，不依赖固定列）
    category_col = None
    main_header_row = None
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            if normalize_text(get_display_value(r, c)) == "分类":
                category_col = c
                main_header_row = r
                break
        if category_col is not None:
            break

    if category_col is None or main_header_row is None:
        raise ValueError("未找到「分类」表头")

    # 找到 D/M/L 标记行
    dl_row = None
    for r in range(main_header_row, min(main_header_row + 5, ws.max_row) + 1):
        markers = 0
        for c in range(1, ws.max_column + 1):
            if normalize_text(get_display_value(r, c)) in {"D", "M", "L"}:
                markers += 1
        if markers >= 2:
            dl_row = r
            break

    if dl_row is None:
        raise ValueError("未找到 D/M/L 标记行")

    header_rows = list(range(main_header_row, dl_row + 1))

    def get_header_texts(col: int) -> list[str]:
        texts = []
        for row in header_rows:
            text = normalize_text(get_display_value(row, col))
            if text:
                texts.append(text)
        return texts

    cut_col = None
    trim_group_cols: list[int] = []
    trim_cols: Dict[str, int] = {}
    weight_cols: Dict[str, int] = {}
    size_cols: Dict[str, int] = {}
    hair_length_col = None
    pound_hair_col = None
    density_col = None
    shape_col = None
    beauty_cols: Dict[str, int] = {}
    note_cols: list[int] = []

    for c in range(1, ws.max_column + 1):
        texts = get_header_texts(c)
        if not texts:
            continue

        combined_text = "".join(texts)
        dml_marker = normalize_text(get_display_value(dl_row, c))

        if "裁断" in combined_text:
            cut_col = c
        elif "整毛" in combined_text:
            trim_group_cols.append(c)
        elif "拉尖" in combined_text:
            trim_cols["拉尖"] = c
        elif "对裁" in combined_text:
            trim_cols["对裁"] = c
        elif "重量" in combined_text and dml_marker in {"D", "M", "L"}:
            weight_cols[dml_marker] = c
        elif "尺数" in combined_text and dml_marker in {"D", "M", "L"}:
            size_cols[dml_marker] = c
        elif "毛长" in combined_text:
            hair_length_col = c
        elif "磅发" in combined_text:
            pound_hair_col = c
        elif "密度" in combined_text:
            density_col = c
        elif "形态" in combined_text:
            shape_col = c
        elif "铝管" in combined_text:
            beauty_cols["铝管"] = c
        elif "方向" in combined_text:
            beauty_cols["方向"] = c
        elif "层数" in combined_text:
            beauty_cols["层数"] = c
        elif "备注" in combined_text:
            note_cols.append(c)

    if "拉尖" not in trim_cols and trim_group_cols:
        trim_group_cols = sorted(trim_group_cols)
        trim_cols["拉尖"] = trim_group_cols[0]
        if len(trim_group_cols) > 1:
            trim_cols["对裁"] = trim_group_cols[1]

    if cut_col is None:
        raise ValueError("未找到「裁断」表头")
    if "拉尖" not in trim_cols:
        raise ValueError("未找到「拉尖」表头")
    if hair_length_col is None:
        raise ValueError("未找到「毛长」表头")

    data_start_row = dl_row + 1

    def build_dml_data(row: int, cols: Dict[str, int], ensure_d: bool = False) -> Dict[str, float]:
        data = {}
        for needle in ["D", "M", "L"]:
            col = cols.get(needle)
            if col is None:
                continue
            value = to_num(ws.cell(row, col).value)
            if value is not None:
                data[needle] = value
        if ensure_d and data and "D" not in data:
            data["D"] = 0
        return data

    def collect_cut_rows(start_row: int) -> list[int]:
        rows = [start_row]
        for r in range(start_row + 1, ws.max_row + 1):
            if len(rows) >= 3:
                break
            if to_str(ws.cell(r, category_col).value):
                break
            if to_num(ws.cell(r, cut_col).value) is None:
                break
            rows.append(r)
        return rows

    def collect_notes(rows: list[int]) -> str:
        values = []
        for row in rows:
            for col in note_cols:
                value = to_str(ws.cell(row, col).value)
                if value and value not in values:
                    values.append(value)
        return "；".join(values)

    def build_cut_weight_list(rows: list[int], include_weight: bool) -> list[Dict[str, Any]]:
        items = []
        for row in rows:
            cut_value = to_num(ws.cell(row, cut_col).value)
            if cut_value is None:
                continue
            item: Dict[str, Any] = {"裁断": cut_value}
            if include_weight:
                row_weight = build_dml_data(row, weight_cols, ensure_d=True)
                if row_weight:
                    item["重量g"] = row_weight
            items.append(item)
        return items

    machine_spec_list = []

    # 读取机器规格清单数据
    for r in range(data_start_row, ws.max_row + 1):
        gear_val = to_str(ws.cell(r, category_col).value)

        if not gear_val:
            continue
        gear_num = to_num(gear_val)
        if gear_num is None or not gear_num.is_integer():
            continue
        number = int(gear_num)

        cut_rows = collect_cut_rows(r)
        trim_data = {
            "拉尖": to_num(ws.cell(r, trim_cols["拉尖"]).value),
            "对裁": to_num(ws.cell(r, trim_cols["对裁"]).value)
            if "对裁" in trim_cols
            else None,
        }
        trim_data = {k: v for k, v in trim_data.items() if v is not None}

        size_data = build_dml_data(r, size_cols)
        double_needle = {
            "毛长": to_num(ws.cell(r, hair_length_col).value) or 0,
            "尺数": size_data,
            "密度": to_num(ws.cell(r, density_col).value) or 0,
        }

        beauty = {
            "铝管": to_num(ws.cell(r, beauty_cols["铝管"]).value) or 0,
            "方向": to_str(ws.cell(r, beauty_cols["方向"]).value),
            "层数": to_num(ws.cell(r, beauty_cols["层数"]).value) or 0,
        }

        cut_weight_list = build_cut_weight_list(cut_rows, include_weight=True)

        machine_spec = {
            "档位": str(number),
            "裁断与重量": cut_weight_list,
            "整毛": trim_data,
            "双针": double_needle,
            "形态": to_str(ws.cell(r, shape_col).value) if shape_col else "",
            "美容": beauty,
        }
        notes = collect_notes(cut_rows)
        if notes:
            machine_spec["备注"] = notes
        machine_spec_list.append(machine_spec)

    manual_spec_list = []
    for r in range(data_start_row, ws.max_row + 1):
        gear_val = to_str(ws.cell(r, category_col).value)
        if not gear_val:
            continue
        if not re.match(r"^[HT]\d+$", gear_val):
            continue

        cut_rows = collect_cut_rows(r)
        trim_data = {
            "拉尖": to_num(ws.cell(r, trim_cols["拉尖"]).value),
            "对裁": to_num(ws.cell(r, trim_cols["对裁"]).value)
            if "对裁" in trim_cols
            else None,
        }
        trim_data = {k: v for k, v in trim_data.items() if v is not None}

        # 人工规格也需要 D/M/L 重量
        cut_weight_list = build_cut_weight_list(cut_rows, include_weight=True)

        pound_hair_col = pound_hair_col or hair_length_col
        pound_hair = to_num(ws.cell(r, pound_hair_col).value)

        manual_spec = {
            "档位": gear_val,
            "整毛": trim_data,
            "裁断与重量": cut_weight_list,
            "双针": {
                "毛长": to_num(ws.cell(r, hair_length_col).value) or 0,
                "磅发": pound_hair,
                "密度": to_num(ws.cell(r, density_col).value) or 0,
            },
            "形态": to_str(ws.cell(r, shape_col).value) if shape_col else "",
            "美容": {
                "铝管": to_num(ws.cell(r, beauty_cols["铝管"]).value) or 0
            },
        }
        notes = collect_notes(cut_rows)
        if notes:
            manual_spec["备注"] = notes
        manual_spec_list.append(manual_spec)

    # 查找工艺说明
    process_notes = {
        "作业方法": "",
        "整毛": "",
        "双针": "",
        "美容": "",
        "制帽": "",
        "手织": "",
        "高针": "",
        "完成": "",
        "包装": ""
    }

    # 查找工艺说明区域（通常在底部）
    for r in range(ws.max_row, 1, -1):
        for c in range(1, ws.max_column + 1):
            v = to_str(ws.cell(r, c).value)
            if "作业方法" in v:
                process_notes["作业方法"] = v
            elif "整毛" in v and ("：" in v or ":" in v):
                process_notes["整毛"] = v
            elif "双针" in v and ("：" in v or ":" in v):
                process_notes["双针"] = v
            elif "美容" in v and ("：" in v or ":" in v):
                process_notes["美容"] = v
            elif "制帽" in v and ("：" in v or ":" in v):
                process_notes["制帽"] = v
            elif "手织" in v and ("：" in v or ":" in v):
                process_notes["手织"] = v
            elif "高针" in v and ("：" in v or ":" in v):
                process_notes["高针"] = v
            elif "完成" in v and ("：" in v or ":" in v):
                process_notes["完成"] = v
            elif "包装" in v and ("：" in v or ":" in v):
                process_notes["包装"] = v

    # 如果没有找到带冒号的，找一些关键文本
    if not process_notes["作业方法"]:
        for r in range(1, ws.max_row + 1):
            v = to_str(ws.cell(r, 1).value)
            if "作业方法" in v or "本规格" in v:
                process_notes["作业方法"] = v
                break

    for key, value in list(process_notes.items()):
        process_notes[key] = strip_note_key(key, value)

    analysis = analyze_imported_specs(machine_spec_list, manual_spec_list)
    wig_type = analysis["wig_type"]
    dml_ratio = analysis["dml_ratio"]

    if wig_type == "间色" and dml_ratio:
        for machine_spec in machine_spec_list:
            machine_spec["DML比值"] = dict(dml_ratio)
    else:
        for machine_spec in machine_spec_list:
            machine_spec.pop("DML比值", None)

    raw_material = find_value_by_label("原材料:", "原材料")
    color_code = ""
    for row in range(1, min(8, ws.max_row) + 1):
        for col in range(1, ws.max_column + 1):
            value = to_str(get_display_value(row, col))
            if re.fullmatch(r"[A-Z]{1,4}\d+/\d+#?", value):
                color_code = value
                break
        if color_code:
            break

    cap_no = find_value_by_label("号码", "号码:")
    cap_desc = find_value_by_label("制帽:")
    hat_mark = find_value_by_label("唛头", "唛头:")
    sample_no = find_value_by_label("规格书", "规格书:")
    customer_no = find_value_by_label("客户", "客户:")
    product_name = find_value_by_label("品名", "品名:")
    if product_name == "0":
        product_name = ""
    cap_value = clean_meta_value(cap_desc) or clean_meta_value(cap_no)

    meta = {
        "_id": "",
        "样品编号": clean_meta_value(sample_no),
        "假发类型": wig_type,
        "客户编号": clean_meta_value(customer_no),
        "品名": clean_meta_value(product_name),
        "原材料": clean_meta_value(raw_material),
        "CAP": cap_value,
        "文件名称": os.path.basename(path),
        "颜色编号": clean_meta_value(color_code),
        "发丝种类": "+".join(
            dict.fromkeys(re.findall(r"([A-Za-z]+)\s*:", raw_material))
        ),
        "唛头": clean_meta_value(hat_mark),
    }

    return {
        "元数据": meta,
        "机器规格清单": machine_spec_list,
        "人工规格清单": manual_spec_list,
        "工艺说明": process_notes
    }




def generate_full_json(parsed_data: Dict[str, Any]) -> Dict[str, Any]:
    """生成完整的沐茵丝假发成品稿 JSON，符合当前类型定义。"""
    meta = parsed_data.get("元数据", {})
    machine_spec_list = parsed_data["机器规格清单"]
    manual_spec_list = parsed_data["人工规格清单"]
    process_notes = parsed_data["工艺说明"]

    return {
        "沐茵丝假发成品稿": {
            "_id": meta.get("_id", ""),
            "样品编号": meta.get("样品编号", ""),
            "假发类型": meta.get("假发类型", "纯色"),
            "客户编号": meta.get("客户编号", ""),
            "品名": meta.get("品名", ""),
            "原材料": meta.get("原材料", ""),
            "CAP": meta.get("CAP", ""),
            "tag": "成品稿",
            "文件名称": meta.get("文件名称", ""),
            "染色档位列表": [],
            "制品规格书": {
                "机器规格清单": machine_spec_list,
                "人工规格清单": manual_spec_list,
                "胶丝比例id": {
                    "颜色编号": meta.get("颜色编号", ""),
                    "发丝种类": meta.get("发丝种类", ""),
                },
                "制帽": {
                    "唛头": meta.get("唛头", ""),
                },
                "工艺说明": {
                    "作业方法": process_notes.get("作业方法", ""),
                    "整毛": process_notes.get("整毛", ""),
                    "双针": process_notes.get("双针", ""),
                    "美容": process_notes.get("美容", ""),
                    "制帽": process_notes.get("制帽", ""),
                    "手织": process_notes.get("手织", ""),
                    "高针": process_notes.get("高针", ""),
                    "完成": process_notes.get("完成", ""),
                    "包装": process_notes.get("包装", ""),
                },
                "工程重量": {
                    "整毛": {"加减": 0},
                    "双针": {"加减": 0},
                    "美容": {"加减": 0},
                    "制帽": {"加减": 0},
                    "手织": {"加减": 0},
                    "高针": {"加减": 0},
                    "剪驳": {"加减": 0},
                    "发网": {"加减": 0},
                    "完成": {"加减": 0},
                },
            },
            "高针指示单": {
                "注意事项": "高针 :1.高针后帽子不能变形。",
                "高针图": {
                    "json": "",
                    "车线": [],
                    "标注样式": {},
                    "自动修改器": [],
                },
            },
            "手织指示单": {
                "注意事项": "手织 :1.手织后帽子不能变形。",
                "手织图": {
                    "json": "",
                    "间色比例": {
                        "type": "特殊",
                    },
                },
            },
            "头型图片": [],
        }
    }


def main():
    """主函数，执行完整的导出流程"""
    print("开始解析Excel文件...")
    parsed_data = parse_excel(excel_path)
    print(f"解析完成，获得 {len(parsed_data['机器规格清单'])} 条机器规格")
    print(f"获得 {len(parsed_data['人工规格清单'])} 条人工规格")
    print(f"识别假发类型：{parsed_data['元数据']['假发类型']}")

    # 生成符合JsonType的完整数据
    full_data = generate_full_json(parsed_data)
    
    # 导出完整规格数据
    combined_path = os.path.join(output_dir, "完整规格数据.json")
    with open(combined_path, 'w', encoding='utf-8') as f:
        json.dump(full_data, f, ensure_ascii=False, indent=2)
    print(f"完整规格数据已导出到: {combined_path}")

    print("\n数据导出成功！")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"错误: {e}")
        import traceback
        traceback.print_exc()
